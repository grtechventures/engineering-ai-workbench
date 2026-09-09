"""Released, bounded workbook inspection inside the isolated Python worker."""
import tempfile
from pathlib import Path
from .worker import run_plot_script
INSPECT = '''import json
import openpyxl
files=json.load(open('/inputs/data.json'))['files']
result=[]
for file in files[:10]:
 if not file['name'].lower().endswith('.xlsx'): continue
 with open(file['path'],'rb') as stream:
  wb=openpyxl.load_workbook(stream,read_only=True,data_only=True,keep_links=False)
  sheets=[]
  for ws in wb.worksheets[:5]:
   rows=[]
   for row in ws.iter_rows(min_row=1,max_row=20,max_col=30,values_only=True):
    rows.append([str(v)[:150] if v is not None else None for v in row])
   sheets.append({'sheet':ws.title,'rows_preview':rows,'reported_rows':ws.max_row,'reported_columns':ws.max_column})
  wb.close()
 result.append({'name':file['name'],'sheets':sheets})
json.dump({'workbooks':result,'notice':'First 20 rows only. Cached values; formulas not recalculated. Inspect complete rows in the calculation.'},open('/outputs/result.json','w'))
'''
def preview_excel(engine,data,image):
    with tempfile.TemporaryDirectory(dir=engine.data) as directory:
        _,result,_=run_plot_script(INSPECT,{'files':data['files']},Path(directory)/'preview',image,mounts=engine.upload_mounts(data['files']))
        return result
