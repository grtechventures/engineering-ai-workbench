"""Read-only file readers used by generated engineering tasks."""
import csv,json
from pathlib import Path

def kind(file):
    ext=Path(file['name']).suffix.lower()
    formats={'.xlsx':'excel','.csv':'csv','.tsv':'tsv','.pdf':'pdf','.docx':'word','.txt':'text','.md':'text','.json':'json'}
    if ext not in formats:raise ValueError('No installed reader for '+ext+'. Supported: xlsx, csv, tsv, pdf, docx, txt, md, json')
    return formats[ext]

def rows(file,sheet=None):
    """Yield all raw rows, including headings and totals. Do not sum blindly."""
    k=kind(file)
    if k=='excel':
        import openpyxl
        with open(file['path'],'rb') as stream:
            wb=openpyxl.load_workbook(stream,read_only=True,data_only=True,keep_links=False)
            try:
                ws=wb[sheet] if sheet else wb.active
                for row in ws.iter_rows(values_only=True):yield list(row)
            finally:wb.close()
    elif k in ('csv','tsv'):
        with open(file['path'],encoding='utf-8-sig',newline='') as stream:
            yield from csv.reader(stream,delimiter='\t' if k=='tsv' else ',')
    else:raise ValueError('This file is not tabular; use text_blocks')

def text_blocks(file):
    k=kind(file)
    if k=='pdf':
        from pypdf import PdfReader
        with open(file['path'],'rb') as stream:
            reader=PdfReader(stream)
            if reader.is_encrypted:raise ValueError('Encrypted PDFs require an unprotected copy')
            for i,page in enumerate(reader.pages):yield {'location':'page '+str(i+1),'text':page.extract_text() or ''}
    elif k=='word':
        from docx import Document
        with open(file['path'],'rb') as stream:
            doc=Document(stream)
            for i,p in enumerate(doc.paragraphs):yield {'location':'paragraph '+str(i+1),'text':p.text}
            for i,t in enumerate(doc.tables):
                for j,r in enumerate(t.rows):yield {'location':f'table {i+1} row {j+1}','text':' | '.join(c.text for c in r.cells)}
    elif k in ('text','json'):
        with open(file['path'],encoding='utf-8-sig') as stream:
            while True:
                s=stream.read(4096)
                if not s:break
                yield {'location':'text chunk','text':s}
    else:raise ValueError('Use rows for tabular files')

def inspect(file):
    import itertools
    k=kind(file)
    if k in ('excel','csv','tsv'):
        sheets=[None]
        if k=='excel':
            import openpyxl
            with open(file['path'],'rb') as stream:
                wb=openpyxl.load_workbook(stream,read_only=True,data_only=True,keep_links=False);sheets=wb.sheetnames[:5];wb.close()
        return {'name':file['name'],'kind':k,'sheets':[{'sheet':s,'rows':[{'row':i+1,'values':[str(v)[:150] if v is not None else None for v in r[:30]]} for i,r in enumerate(itertools.islice(rows(file,s),30))]} for s in sheets],'notice':'Preview only. Read full rows for calculations. Excel uses cached values without recalculation.'}
    blocks=[];budget=12000
    for b in itertools.islice(text_blocks(file),100):
        text=b['text'][:budget];blocks.append({'location':b['location'],'text':text});budget-=len(text)
        if budget<=0:break
    return {'name':file['name'],'kind':k,'blocks':blocks,'notice':'Bounded text preview. Scanned PDFs may need OCR, which is not installed.'}
